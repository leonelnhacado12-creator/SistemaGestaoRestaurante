from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r'C:\Users\leone\Documents\Relatorio_Sistema_Gestao_Restaurante.docx'
URL = 'https://github.com/leonelnhacado12-creator/SistemaGestaoRestaurante'
doc = Document()
sec = doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.85); sec.right_margin=Inches(.85)
styles=doc.styles; styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10.5); styles['Normal'].paragraph_format.space_after=Pt(6)
for s in ['Title','Heading 1','Heading 2']:
    styles[s].font.name='Aptos Display'; styles[s].font.color.rgb=RGBColor(0,0,0)
styles['Title'].font.size=Pt(25); styles['Heading 1'].font.size=Pt(17); styles['Heading 2'].font.size=Pt(13)
def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def table(headers, rows):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,'173B5F'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row): cells[i].text=str(val); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if ri%2==1:
            for c in cells: shade(c,'EAF1F7')
    doc.add_paragraph()
    return t
def bullets(items):
    for x in items: doc.add_paragraph(x, style='List Bullet')

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('RELATÓRIO DO PROJECTO').bold=True
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Sistema de Gestão de Restaurante')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Análise, requisitos e arquitectura da solução').italic=True
doc.add_paragraph(); doc.add_paragraph('Repositório GitHub: '+URL, style='Subtitle'); doc.add_paragraph('Data: 19 de Outubro de 2026'); doc.add_page_break()
doc.add_heading('Índice', level=1)
toc = doc.add_paragraph()
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'TOC \\o "1-2" \\h \\z \\u')
run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = 'Actualize este índice no Word com botão direito e Actualizar campo.'; run.append(text); fld.append(run); toc._p.append(fld)
doc.add_page_break()

doc.add_heading('1 Introdução', level=1)
doc.add_paragraph('Este relatório apresenta o Sistema de Gestão de Restaurante, uma aplicação desktop desenvolvida em Java para apoiar as operações de atendimento, vendas, pedidos, clientes, produtos, promoções, inventário, utilizadores e relatórios. O projecto centraliza as actividades do restaurante numa aplicação única, com persistência em MySQL, controlo de acesso por perfis e geração de recibos em PDF.')
doc.add_paragraph('O trabalho foi organizado segundo uma arquitectura em camadas, separando interface, serviços, regras de negócio, acesso a dados e modelos. Esta abordagem facilita a manutenção, os testes e a evolução futura do sistema.')

doc.add_heading('2 Tema e descrição do projecto', level=1)
doc.add_paragraph('O tema do projecto é a informatização da gestão operacional de um restaurante. A solução permite registar produtos e preços, receber pedidos, aplicar promoções, gerir clientes, acompanhar stock, processar pagamentos em dinheiro, cartão, M-Pesa, e-Mola e mKesh, gerar recibos PDF e consultar indicadores de actividade.')
doc.add_paragraph('A aplicação contempla três perfis de acesso: Administrador, Gerente e Atendente. Cada perfil visualiza apenas as operações necessárias à sua função. A base de dados MySQL mantém os dados de forma persistente e o sistema cria automaticamente as tabelas necessárias no arranque.')

doc.add_heading('3 Objectivos', level=1)
bullets(['Digitalizar o ciclo de venda, desde a selecção do produto até ao recibo.', 'Reduzir erros no registo de preços, quantidades, stock e pagamentos.', 'Controlar o acesso às funções de acordo com o perfil do funcionário.', 'Disponibilizar informação para apoio à decisão através do dashboard e relatórios.', 'Criar uma base extensível para mesas, reservas, fornecedores, lotes e auditoria.'])

doc.add_heading('4 Requisitos funcionais', level=1)
table(['ID','Requisito','Descrição'],[
('RF01','Autenticação','Permitir iniciar e terminar sessão com utilizador e senha protegida por hash.'),('RF02','Perfis de acesso','Aplicar permissões distintas para Administrador, Gerente e Atendente.'),('RF03','CRUD de produtos','Criar, consultar, editar, remover, pesquisar e ordenar produtos.'),('RF04','CRUD de clientes','Registar, editar, remover e pesquisar clientes.'),('RF05','Gestão de pedidos','Criar pedidos, adicionar/remover itens, calcular subtotal, desconto e total.'),('RF06','Pagamentos','Registar dinheiro, cartão, M-Pesa, e-Mola e mKesh, incluindo telefone móvel.'),('RF07','Pedidos históricos','Consultar pedidos, filtrar por estado e avançar o estado operacional.'),('RF08','Promoções','Criar, editar, activar, desactivar e aplicar descontos.'),('RF09','Inventário','Gerir itens, entradas, saídas, stock mínimo, fornecedores, lotes e validade.'),('RF10','Mesas e reservas','Registar mesas, estados, reservas, clientes, horários e número de pessoas.'),('RF11','Recibos','Gerar e descarregar recibos PDF dos pedidos.'),('RF12','Relatórios','Exportar produtos, promoções e pedidos em CSV.'),('RF13','Dashboard','Apresentar produtos, pedidos, facturação, alertas de stock e pagamentos.'),('RF14','Auditoria','Registar operação, utilizador, entidade, detalhe e data/hora.')])

doc.add_heading('5 Requisitos não funcionais', level=1)
table(['ID','Categoria','Requisito'],[
('RNF01','Usabilidade','Interface consistente, clara, com filtros, tabelas ordenáveis e mensagens de validação.'),('RNF02','Manutenibilidade','Separação em pacotes model, dao, service, controller, ui, security e report.'),('RNF03','Persistência','Utilizar MySQL e transacções para garantir consistência dos pedidos.'),('RNF04','Segurança','Nunca guardar senhas em texto simples e restringir operações por perfil.'),('RNF05','Desempenho','Usar índices, consultas filtradas e carregamento adequado das tabelas.'),('RNF06','Portabilidade','Executar em Java 17 ou superior, com Maven e NetBeans.'),('RNF07','Disponibilidade','Apresentar o estado da ligação à base de dados e mensagens de erro compreensíveis.'),('RNF08','Integridade','Validar campos obrigatórios, preços, quantidades, descontos, telefones e estados.'),('RNF09','Evolução','Permitir adicionar módulos sem misturar interface com acesso à base de dados.'),('RNF10','Exportação','Gerar ficheiros CSV em UTF-8 e recibos PDF legíveis.')])

doc.add_heading('6 Arquitectura e diagrama de classes', level=1)
doc.add_paragraph('A arquitectura segue MVC/camadas. A interface Swing comunica com o AplicacaoController; os serviços aplicam as regras de negócio; os DAOs executam operações MySQL; os modelos representam o domínio.')
diagram='''+----------------------+        +----------------------+
| AplicacaoController |------->| DadosSistema        |
| sessão e permissões |        | serviços             |
+----------+-----------+        +----+-----------+-----+
           |                         |           |
           v                         v           v
      +---------+              +----------+  +---------+
      | Telas   |              | Services |  | Reports |
      | Swing   |              +----+-----+  +---------+
      +---------+                   |
                                    v
                              +-----+-----+
                              | DAO       |
                              +-----+-----+
                                    |
                                    v
                              +-----+-----+
                              | MySQL     |
                              +-----------+

Usuario <|-- Gerente
Pagamento <|.. PagamentoDinheiro
Pagamento <|.. PagamentoCartao
Pagamento <|.. PagamentoMPesa
Pagamento <|.. PagamentoEMola
Pagamento <|.. PagamentoMKesh
Pedido o-- ItemPedido
Pedido --> Cliente
ItemPedido --> Produto
LoteStock --> Fornecedor
Reserva --> Mesa'''
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; run=p.add_run(diagram); run.font.name='Consolas'; run.font.size=Pt(8.5)

doc.add_heading('7 Tecnologias e organização', level=1)
bullets(['Java 17+, Swing e NetBeans GUI Builder para a interface.', 'Maven para dependências e execução.', 'MySQL e JDBC para persistência.', 'PDF e CSV para recibos e relatórios.', 'Git e GitHub para versionamento e entrega.'])
doc.add_paragraph('Repositório: '+URL)

doc.add_heading('8 Conclusões', level=1)
doc.add_paragraph('O Sistema de Gestão de Restaurante responde às necessidades essenciais de uma operação de restauração, centralizando pedidos, produtos, clientes, pagamentos, inventário, relatórios e controlo de acesso. A arquitectura em camadas facilita a manutenção e a evolução, enquanto a persistência MySQL evita a perda de informação entre execuções.')
doc.add_paragraph('Como evolução, recomenda-se concluir a ligação persistente dos módulos de mesas, reservas, fornecedores e lotes, adicionar testes de integração com MySQL e configurar backups automáticos. Estas melhorias podem ser incorporadas sem alterar a separação principal do projecto.')
doc.add_heading('9 Execução e entrega', level=1)
doc.add_paragraph('Para executar o sistema, é necessário Java 17 ou superior, Maven, MySQL e a base gestao_restaurante. O script SQL está no directório database e o README do repositório contém a configuração das variáveis de conexão. O código fonte e a documentação estão disponíveis no repositório GitHub indicado neste relatório.')

doc.save(OUT)
print(OUT)
